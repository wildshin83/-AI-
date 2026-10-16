"""확정 Inventor 업로드를 BOM Rev·품목 스냅샷으로 전환한다."""

from __future__ import annotations

from decimal import Decimal, InvalidOperation
from io import BytesIO
from pathlib import Path

from sqlalchemy import select
from sqlalchemy.orm import Session

from app.domain.bom_compare import BomChange, BomComparableItem, compare_bom_revisions
from app.domain.part_matching import normalize_drawing_number
from app.models.bom import BomItem, BomRevision
from app.models.imports import UploadBatch
from app.services.audit import record_audit


def _text(value: object) -> str | None:
    cleaned = str(value).strip() if value is not None else ""
    return cleaned or None


def _quantity(value: object) -> Decimal:
    try:
        return Decimal(str(value))
    except (InvalidOperation, TypeError) as error:
        raise ValueError(f"수량 값 '{value}'을(를) 읽을 수 없습니다.") from error


def create_bom_revision(
    session: Session,
    *,
    project_id: int,
    upload_batch: UploadBatch,
    revision_name: str,
    memo: str | None,
) -> BomRevision:
    """확정된 Inventor 업로드 한 건을 변경 불가 BOM Rev로 저장한다."""

    from openpyxl import load_workbook

    if upload_batch.source_type != "inventor_bom" or upload_batch.status != "confirmed":
        raise ValueError("확정된 Inventor BOM 업로드만 BOM Rev로 등록할 수 있습니다.")
    if not revision_name.strip():
        raise ValueError("BOM Rev 이름은 필수입니다.")
    workbook_source = (
        BytesIO(upload_batch.file_content)
        if upload_batch.file_content is not None
        else Path(upload_batch.stored_path)
    )
    workbook = load_workbook(workbook_source, read_only=True, data_only=True)
    worksheet = workbook[upload_batch.sheet_name] if upload_batch.sheet_name else workbook.active
    headers = [
        cell.value
        for cell in next(
            worksheet.iter_rows(min_row=upload_batch.header_row, max_row=upload_batch.header_row)
        )
    ]
    positions = {
        str(header).strip(): index for index, header in enumerate(headers) if header is not None
    }
    revision = BomRevision(
        project_id=project_id,
        upload_batch_id=upload_batch.id,
        revision_name=revision_name.strip(),
        memo=memo.strip() if memo else None,
    )
    session.add(revision)
    session.flush()
    for row in worksheet.iter_rows(min_row=(upload_batch.header_row or 1) + 1, values_only=True):

        def value(field: str) -> object:
            header = upload_batch.mapping.get(field)
            index = positions.get(header) if header else None
            return row[index] if index is not None and index < len(row) else None

        drawing_number = _text(value("drawing_number"))
        if not drawing_number:
            continue
        session.add(
            BomItem(
                bom_revision_id=revision.id,
                source_item_number=_text(value("item_number")),
                drawing_number=drawing_number,
                normalized_drawing_number=normalize_drawing_number(drawing_number),
                title=_text(value("title")),
                component=_text(value("component")),
                material=_text(value("material")),
                treatment=_text(value("treatment")),
                specification=_text(value("specification")),
                maker=_text(value("maker")),
                quantity=_quantity(value("quantity")),
                source_data={},
            )
        )
    record_audit(
        session,
        action="bom_revision.created",
        entity_type="bom_revision",
        entity_id=revision.id,
        details={"project_id": project_id, "revision_name": revision.revision_name},
    )
    session.commit()
    session.refresh(revision)
    return revision


def revision_changes(session: Session, revision: BomRevision) -> list[BomChange]:
    """같은 프로젝트의 직전 Rev와 현재 Rev의 차이를 계산한다."""

    previous = session.scalar(
        select(BomRevision)
        .where(
            BomRevision.project_id == revision.project_id,
            BomRevision.created_at < revision.created_at,
        )
        .order_by(BomRevision.created_at.desc())
    )
    if previous is None:
        return []

    def comparable(revision_id: int) -> list[BomComparableItem]:
        return [
            BomComparableItem(
                drawing_number=item.drawing_number,
                normalized_drawing_number=item.normalized_drawing_number,
                material=item.material,
                treatment=item.treatment,
                specification=item.specification,
                maker=item.maker,
                quantity=item.quantity,
            )
            for item in session.scalars(
                select(BomItem).where(BomItem.bom_revision_id == revision_id)
            )
        ]

    return compare_bom_revisions(comparable(previous.id), comparable(revision.id))
