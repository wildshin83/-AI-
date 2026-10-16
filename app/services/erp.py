"""ERP Excel 이력 적재와 BOM 후보 매칭 유스케이스."""

from __future__ import annotations

from dataclasses import dataclass
from datetime import date, datetime, time, timedelta
from decimal import Decimal, InvalidOperation
from io import BytesIO
from pathlib import Path

from openpyxl import load_workbook
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.domain.part_matching import normalize_drawing_number, normalize_specification
from app.models.alerts import PriceRecord
from app.models.bom import BomItem
from app.models.erp import ErpMatch, ERPTransaction
from app.models.imports import UploadBatch
from app.services.audit import record_audit


class ErpImportAlreadyExistsError(ValueError):
    """이미 이력으로 적재한 ERP 업로드를 다시 선택했을 때 발생한다."""


@dataclass(frozen=True)
class ErpCandidate:
    transaction: ERPTransaction
    match_method: str


def _text(value: object) -> str | None:
    cleaned = str(value).strip() if value is not None else ""
    return cleaned or None


def _decimal(value: object) -> Decimal | None:
    text = _text(value)
    if text is None:
        return None
    try:
        return Decimal(text.replace(",", ""))
    except InvalidOperation as error:
        raise ValueError(f"숫자 값 '{value}'을(를) 읽을 수 없습니다.") from error


def _date(value: object) -> date | None:
    if value is None or value == "":
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    try:
        return datetime.fromisoformat(str(value).strip()).date()
    except ValueError as error:
        raise ValueError(f"날짜 값 '{value}'을(를) 읽을 수 없습니다.") from error


def _datetime(value: object) -> datetime | None:
    if value is None or value == "":
        return None
    if isinstance(value, datetime):
        return value
    if isinstance(value, date):
        return datetime.combine(value, time.min)
    try:
        return datetime.fromisoformat(str(value).strip())
    except ValueError as error:
        raise ValueError(f"ERP 기준 일시 '{value}'을(를) 읽을 수 없습니다.") from error


def import_erp_transactions(session: Session, *, upload_batch: UploadBatch) -> int:
    """확정 ERP 업로드를 원본 행 단위 이력으로 적재한다.

    현재 제공된 양식의 ``최종수정일``은 정렬 기준만 되며, 발주일이나 입고일로
    바꾸어 저장하지 않는다.
    """

    if upload_batch.source_type != "erp_bom" or upload_batch.status != "confirmed":
        raise ValueError("확정된 ERP BOM 업로드만 ERP 이력으로 등록할 수 있습니다.")
    already_imported = session.scalar(
        select(ERPTransaction.id).where(ERPTransaction.upload_batch_id == upload_batch.id)
    )
    if already_imported is not None:
        raise ErpImportAlreadyExistsError("선택한 ERP 업로드는 이미 이력으로 등록되어 있습니다.")

    workbook_source = (
        BytesIO(upload_batch.file_content)
        if upload_batch.file_content is not None
        else Path(upload_batch.stored_path)
    )
    workbook = load_workbook(workbook_source, read_only=True, data_only=True)
    worksheet = workbook[upload_batch.sheet_name] if upload_batch.sheet_name else workbook.active
    header_row = upload_batch.header_row or 1
    headers = [
        cell.value for cell in next(worksheet.iter_rows(min_row=header_row, max_row=header_row))
    ]
    positions = {
        str(header).strip(): index for index, header in enumerate(headers) if header is not None
    }
    count = 0
    for row_number, row in enumerate(
        worksheet.iter_rows(min_row=header_row + 1, values_only=True), start=header_row + 1
    ):

        def value(field: str) -> object:
            header = upload_batch.mapping.get(field)
            position = positions.get(header) if header else None
            return row[position] if position is not None and position < len(row) else None

        part_number = _text(value("part_number"))
        if not part_number:
            continue
        specification = _text(value("specification"))
        source_data = {
            field: _text(value(field))
            for field in upload_batch.mapping
            if _text(value(field)) is not None
        }
        try:
            session.add(
                ERPTransaction(
                    project_id=upload_batch.project_id,
                    upload_batch_id=upload_batch.id,
                    part_number=part_number,
                    normalized_part_number=normalize_drawing_number(part_number),
                    part_name=_text(value("part_name")),
                    material=_text(value("material")),
                    specification=specification,
                    normalized_specification=normalize_specification(specification),
                    maker=_text(value("maker")),
                    supplier=_text(value("supplier")),
                    item_category=_text(value("item_category")),
                    unit=_text(value("unit")),
                    unit_price=_decimal(value("unit_price")),
                    ordered_quantity=_decimal(value("ordered_quantity")),
                    received_quantity=_decimal(value("received_quantity")),
                    delivered_quantity=_decimal(value("delivered_quantity")),
                    required_quantity=_decimal(value("required_quantity")),
                    planned_delivery_date=_date(value("planned_delivery_date")),
                    revised_delivery_date=_date(value("revised_delivery_date")),
                    erp_updated_at=_datetime(value("erp_updated_at")),
                    purchase_order_number=_text(value("purchase_order_number")),
                    eco_number=_text(value("eco_number")),
                    erp_version=_text(value("erp_version")),
                    source_data={"row": row_number, "values": source_data},
                )
            )
        except ValueError as error:
            raise ValueError(f"ERP Excel {row_number}행: {error}") from error
        count += 1
    record_audit(
        session,
        action="erp_transactions.imported",
        entity_type="upload_batch",
        entity_id=upload_batch.id,
        details={"project_id": upload_batch.project_id, "row_count": count},
    )
    session.commit()
    return count


def _recent_transactions(session: Session, *, project_id: int, as_of: date) -> list[ERPTransaction]:
    cutoff = datetime.combine(as_of - timedelta(days=365), time.min)
    return list(
        session.scalars(
            select(ERPTransaction)
            .where(
                ERPTransaction.project_id == project_id,
                ERPTransaction.erp_updated_at >= cutoff,
            )
            .order_by(ERPTransaction.erp_updated_at.desc(), ERPTransaction.id.desc())
        )
    )


def find_erp_candidates(
    session: Session, *, bom_item: BomItem, as_of: date | None = None
) -> list[ErpCandidate]:
    """최근 1년 ERP 이력을 도면번호 우선, 규격 보조 순서로 반환한다."""

    reference_date = as_of or date.today()
    drawing_number = normalize_drawing_number(bom_item.drawing_number)
    drawing_matches = _recent_transactions(
        session, project_id=_project_id_for_item(session, bom_item), as_of=reference_date
    )
    exact = [
        transaction
        for transaction in drawing_matches
        if transaction.normalized_part_number == drawing_number
    ]
    if exact:
        return [ErpCandidate(transaction, "drawing_number") for transaction in exact]

    specification = normalize_specification(bom_item.specification)
    if not specification:
        return []
    return [
        ErpCandidate(transaction, "specification")
        for transaction in drawing_matches
        if transaction.normalized_specification == specification
    ]


def _project_id_for_item(session: Session, bom_item: BomItem) -> int:
    from app.models.bom import BomRevision

    project_id = session.scalar(
        select(BomRevision.project_id).where(BomRevision.id == bom_item.bom_revision_id)
    )
    if project_id is None:
        raise ValueError("BOM 품목의 프로젝트를 찾을 수 없습니다.")
    return project_id


def confirm_erp_match(
    session: Session,
    *,
    bom_item: BomItem,
    transaction: ERPTransaction,
    match_method: str,
    review_type: str,
    review_evidence: dict[str, str],
    confirmation_note: str,
    confirmed_by: str,
) -> ErpMatch:
    """사용자가 검토 근거를 남긴 ERP 후보만 적용 상태로 확정한다."""

    if transaction.project_id != _project_id_for_item(session, bom_item):
        raise ValueError("같은 프로젝트의 ERP 이력만 연결할 수 있습니다.")
    if match_method not in {"drawing_number", "specification", "manual"}:
        raise ValueError("알 수 없는 매칭 방식입니다.")
    if review_type not in {"machined", "purchased"}:
        raise ValueError("가공품 또는 구매품 검토 유형을 선택하세요.")
    required_evidence = (
        ("shape_dimension_checked", "material_checked", "treatment_checked", "quantity_checked")
        if review_type == "machined"
        else ("specification_checked", "maker_checked")
    )
    missing = [key for key in required_evidence if not review_evidence.get(key)]
    if missing:
        raise ValueError("품목 유형에 필요한 확인 항목을 모두 기록하세요.")
    if not confirmation_note.strip() or not confirmed_by.strip():
        raise ValueError("확정 근거와 확정자는 필수입니다.")

    session.query(ErpMatch).filter(
        ErpMatch.bom_item_id == bom_item.id, ErpMatch.status == "applied"
    ).update({"status": "held"}, synchronize_session=False)
    match = ErpMatch(
        bom_item_id=bom_item.id,
        erp_transaction_id=transaction.id,
        match_method=match_method,
        status="applied",
        review_type=review_type,
        review_evidence=review_evidence,
        confirmation_note=confirmation_note.strip(),
        confirmed_by=confirmed_by.strip(),
        confirmed_at=datetime.now(),
    )
    session.add(match)
    if transaction.unit_price is not None:
        session.add(
            PriceRecord(
                bom_item_id=bom_item.id,
                erp_transaction_id=transaction.id,
                price_type="erp_reference",
                amount=transaction.unit_price,
                basis_date=(
                    transaction.erp_updated_at.date() if transaction.erp_updated_at else None
                ),
                memo="사용자 확정 ERP 후보의 참고 단가",
            )
        )
    if bom_item.item_type is None:
        bom_item.item_type = review_type
    session.flush()
    record_audit(
        session,
        action="erp_match.confirmed",
        entity_type="erp_match",
        entity_id=match.id,
        details={
            "bom_item_id": bom_item.id,
            "erp_transaction_id": transaction.id,
            "match_method": match_method,
            "review_type": review_type,
        },
    )
    session.commit()
    session.refresh(match)
    return match
