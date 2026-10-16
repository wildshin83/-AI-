"""BOM·구매 협의·고객 견적 Excel을 위한 공통 출력 함수."""

from __future__ import annotations

from pathlib import Path
from re import match

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font

from app.domain.bom_compare import BomChange
from app.models.alerts import Alert
from app.models.bom import BomItem
from app.models.quotes import QuoteItem, QuoteVersion

CUSTOMER_FORBIDDEN_KEYS = {"supplier", "erp_price", "internal_price", "purchase_price"}
CustomerQuoteLine = tuple[str, str, str, float, float]


def _customer_template_slots(sheet) -> dict[str, list[int]]:
    """LGES 2번 시트의 장비·부품 입력 행만 안전하게 식별한다."""

    slots = {"purchased": [], "machined": []}
    section: str | None = None
    for row_index in range(5, 54):
        section_name = sheet.cell(row_index, 2).value
        if section_name in {"Purchasing Item", "Processing Item"}:
            section = section_name
        sequence = str(sheet.cell(row_index, 3).value or "")
        if not match(r"^\d+\)", sequence):
            continue
        if section == "Purchasing Item":
            slots["purchased"].append(row_index)
        elif section == "Processing Item":
            slots["machined"].append(row_index)
    return slots


def _customer_quote_lines(items: list[QuoteItem]) -> dict[str, list[CustomerQuoteLine]]:
    """판매 견적 스냅샷을 고객 양식의 구매/가공 행으로 전개한다."""

    lines = {"purchased": [], "machined": []}
    for item in items:
        quantity = float(item.quantity)
        item_name = item.item_name or item.drawing_number
        if item.item_type == "machined":
            components = (
                ("Raw Material", item.material or item.drawing_number, item.material_cost),
                ("Processing Cost", item.drawing_number, item.machining_cost),
                ("Post Processing", item.treatment or item.drawing_number, item.treatment_cost),
            )
            for label, detail, unit_price in components:
                if unit_price:
                    lines["machined"].append(
                        (f"{item_name} - {label}", detail, "-", quantity, float(unit_price))
                    )
        else:
            lines["purchased"].append(
                (
                    item_name,
                    item.specification or item.drawing_number,
                    item.maker or "-",
                    quantity,
                    float(item.purchased_unit_price),
                )
            )
    return lines


def _header(sheet, values: list[str]) -> None:
    sheet.append(values)
    for cell in sheet[1]:
        cell.font = Font(name="Arial", bold=True)


def export_bom(path: Path, *, items: list[BomItem], alerts: list[Alert]) -> Path:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "BOM"
    _header(sheet, ["도면번호", "품명", "구분", "재질", "후처리", "규격", "제조사", "수량", "경고"])
    alert_numbers = {
        alert.bom_item_id: alert.title for alert in alerts if alert.status != "resolved"
    }
    for item in items:
        sheet.append(
            [
                item.drawing_number,
                item.component or item.title,
                item.item_type,
                item.material,
                item.treatment,
                item.specification,
                item.maker,
                float(item.quantity),
                alert_numbers.get(item.id, ""),
            ]
        )
    workbook.save(path)
    return path


def export_purchase_coordination(path: Path, *, items: list[BomItem], alerts: list[Alert]) -> Path:
    """구매팀 협의용 사양·위험·조치 상태 목록을 생성한다."""

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "구매 협의"
    _header(sheet, ["도면번호", "품명", "규격", "제조사", "수량", "경고", "상태"])
    alerts_by_item = {alert.bom_item_id: alert for alert in alerts if alert.status != "resolved"}
    for item in items:
        alert = alerts_by_item.get(item.id)
        sheet.append(
            [
                item.drawing_number,
                item.component or item.title,
                item.specification,
                item.maker,
                float(item.quantity),
                alert.title if alert else "",
                alert.status if alert else "",
            ]
        )
    workbook.save(path)
    return path


def export_bom_revision_comparison(path: Path, *, changes: list[BomChange]) -> Path:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "BOM Rev 비교"
    _header(sheet, ["도면번호", "변경 유형", "변경 항목"])
    for change in changes:
        sheet.append([change.drawing_number, change.change_type, ", ".join(change.changed_fields)])
    workbook.save(path)
    return path


def export_cost_summary(path: Path, *, items: list[QuoteItem]) -> Path:
    """고객 판매 견적 구성비를 내부 검토용으로 집계한다."""

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "원가 집계"
    _header(
        sheet,
        ["도면번호", "품명", "수량", "재질비", "가공비", "후처리비", "구매품 단가", "판매금액"],
    )
    for item in items:
        unit_price = (
            item.material_cost + item.machining_cost + item.treatment_cost
            if item.item_type == "machined"
            else item.purchased_unit_price
        )
        sheet.append(
            [
                item.drawing_number,
                item.item_name,
                float(item.quantity),
                float(item.material_cost),
                float(item.machining_cost),
                float(item.treatment_cost),
                float(item.purchased_unit_price),
                float(unit_price * item.quantity),
            ]
        )
    workbook.save(path)
    return path


def assert_customer_safe(rows: list[dict[str, object]]) -> None:
    for row in rows:
        forbidden = CUSTOMER_FORBIDDEN_KEYS & row.keys()
        if forbidden:
            raise ValueError(
                f"고객용 출력에 금지 필드가 포함되었습니다: {', '.join(sorted(forbidden))}"
            )


def export_customer_quote(
    template_path: Path, output_path: Path, *, quote: QuoteVersion, items: list[QuoteItem]
) -> Path:
    """고객 양식 2번 시트를 보존한 복사본의 제출 견적 영역에 판매값만 기록한다."""

    rows = [
        {"drawing_number": item.drawing_number, "name": item.item_name, "quantity": item.quantity}
        for item in items
    ]
    assert_customer_safe(rows)
    workbook = load_workbook(template_path)
    sheet = workbook["2.Detail Quotation(Global)"]
    sheet["A2"] = f"Quotation {quote.version_number}"
    slots = _customer_template_slots(sheet)
    lines = _customer_quote_lines(items)
    for item_type, quote_lines in lines.items():
        if len(quote_lines) > len(slots[item_type]):
            raise ValueError(
                f"고객 견적 양식의 {item_type} 입력 가능 행을 초과했습니다. "
                "양식을 확장하거나 품목을 분리해 다시 생성하세요."
            )
        for row_index, (name, specification, maker, quantity, unit_price) in zip(
            slots[item_type], quote_lines
        ):
            sheet.cell(row_index, 4).value = name
            sheet.cell(row_index, 5).value = specification
            sheet.cell(row_index, 6).value = maker
            sheet.cell(row_index, 7).value = quantity
            sheet.cell(row_index, 8).value = "EA"
            sheet.cell(row_index, 9).value = unit_price
    workbook.calculation.calcMode = "auto"
    workbook.calculation.fullCalcOnLoad = True
    workbook.calculation.forceFullCalc = True
    workbook.save(output_path)
    return output_path
