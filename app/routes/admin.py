"""관리자 전용 계정 생성과 Excel 일괄 등록."""

from __future__ import annotations

import os
from datetime import datetime
from io import BytesIO

from fastapi import APIRouter, Depends, File, Form, Request, UploadFile, status
from fastapi.responses import StreamingResponse
from fastapi.templating import Jinja2Templates
from sqlalchemy import select
from sqlalchemy.exc import IntegrityError
from sqlalchemy.orm import Session

from app.db import get_session
from app.models.auth import AccountRequest, AuthUser
from app.services.auth import create_user, normalize_email

router = APIRouter(prefix="/admin", tags=["admin"])
templates = Jinja2Templates(directory="app/templates")
MAX_WORKBOOK_BYTES = 2 * 1024 * 1024
MAX_ACCOUNTS_PER_UPLOAD = 500
HEADER_ALIASES = {
    "email": {"email", "e-mail", "이메일", "사용자이메일", "계정이메일"},
    "name": {"name", "displayname", "사용자명", "이름", "성명"},
    "password": {
        "temporarypassword",
        "temporary_password",
        "temp_password",
        "임시비밀번호",
        "초기비밀번호",
    },
}


@router.get("/users/template")
def account_template() -> StreamingResponse:
    """계정 업로드용 빈 Excel 서식을 제공한다."""

    from openpyxl import Workbook

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "계정"
    sheet.append(["email", "name", "temporary_password"])
    sheet.freeze_panes = "A2"
    sheet.column_dimensions["A"].width = 32
    sheet.column_dimensions["B"].width = 24
    sheet.column_dimensions["C"].width = 28
    buffer = BytesIO()
    workbook.save(buffer)
    buffer.seek(0)
    return StreamingResponse(
        buffer,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": 'attachment; filename="account_template.xlsx"'},
    )


def _clean_header(value: object) -> str:
    return "".join(str(value or "").strip().casefold().replace(" ", "").replace("-", "").split())


def _parse_workbook(contents: bytes, session: Session) -> tuple[list[dict[str, str]], list[str]]:
    from openpyxl import load_workbook

    try:
        workbook = load_workbook(BytesIO(contents), read_only=True, data_only=True)
        sheet = workbook.active
    except Exception as error:
        return [], [f"Excel 파일을 읽을 수 없습니다: {type(error).__name__}"]

    rows = sheet.iter_rows(values_only=True)
    first_row = next(rows, None)
    if first_row is None:
        workbook.close()
        return [], ["Excel 파일에 헤더 행이 없습니다."]
    headers = [_clean_header(value) for value in first_row]
    positions: dict[str, int] = {}
    for field, aliases in HEADER_ALIASES.items():
        normalized_aliases = {_clean_header(alias) for alias in aliases}
        position = next(
            (i for i, header in enumerate(headers) if header in normalized_aliases), None
        )
        if position is not None:
            positions[field] = position

    missing = set(HEADER_ALIASES) - positions.keys()
    if missing:
        workbook.close()
        return [], ["필수 열이 없습니다: email, name, temporary_password"]

    accounts: list[dict[str, str]] = []
    errors: list[str] = []
    seen: set[str] = set()
    existing = set(session.scalars(select(AuthUser.email)))
    for row_number, row in enumerate(rows, start=2):
        if not any(value not in (None, "") for value in row):
            continue
        if len(accounts) >= MAX_ACCOUNTS_PER_UPLOAD:
            errors.append(
                f"최대 {MAX_ACCOUNTS_PER_UPLOAD}개 계정까지만 한 번에 등록할 수 있습니다."
            )
            break
        try:
            raw_email = str(row[positions["email"]] or "")
            email = normalize_email(raw_email)
            display_name = str(row[positions["name"]] or "").strip()
            temporary_password = str(row[positions["password"]] or "")
            if not display_name or len(display_name) > 120:
                raise ValueError("이름이 비어 있거나 120자를 초과합니다.")
            if len(temporary_password) < 12 or len(temporary_password) > 256:
                raise ValueError("임시 비밀번호는 12~256자여야 합니다.")
            if email in seen:
                raise ValueError("파일 안에 이메일이 중복됩니다.")
            if email in existing:
                raise ValueError("이미 등록된 이메일입니다.")
            seen.add(email)
            accounts.append(
                {
                    "email": email,
                    "display_name": display_name,
                    "temporary_password": temporary_password,
                }
            )
        except (ValueError, IndexError) as error:
            errors.append(f"{row_number}행: {error}")
    workbook.close()
    if not accounts:
        errors.append("등록할 계정이 없습니다.")
    return accounts, errors


@router.get("/users")
def user_list(
    request: Request,
    session: Session = Depends(get_session),
):
    users = list(
        session.scalars(select(AuthUser).order_by(AuthUser.created_at.desc(), AuthUser.id.desc()))
    )
    requests = list(
        session.scalars(
            select(AccountRequest)
            .where(AccountRequest.status == "pending")
            .order_by(AccountRequest.created_at, AccountRequest.id)
        )
    )
    return templates.TemplateResponse(
        request,
        "admin/users.html",
        {
            "users": users,
            "requests": requests,
            "message": None,
            "errors": [],
            "single_error": None,
            "smtp_configured": bool(
                os.getenv("BOM_SMTP_HOST")
                and os.getenv("BOM_SMTP_FROM")
                and (os.getenv("BOM_ADMIN_NOTIFICATION_EMAIL") or users)
            ),
        },
    )


@router.post("/account-requests/{request_id}/approve")
def approve_account_request(
    request: Request,
    request_id: int,
    temporary_password: str = Form(),
    session: Session = Depends(get_session),
):
    account_request = session.get(AccountRequest, request_id)
    if not account_request or account_request.status != "pending":
        message, single_error = None, "대기 중인 신청을 찾을 수 없습니다."
    else:
        try:
            user = create_user(
                session,
                email=account_request.email,
                display_name=account_request.display_name,
                temporary_password=temporary_password,
            )
            account_request.status = "approved"
            account_request.reviewed_at = datetime.utcnow()
            session.commit()
            message = (
                f"{user.email} 계정을 승인하고 등록했습니다. "
                "첫 로그인 때 비밀번호 변경이 필요합니다."
            )
            single_error = None
        except ValueError as error:
            session.rollback()
            message, single_error = None, str(error)
        except IntegrityError:
            session.rollback()
            message, single_error = None, "이미 등록된 이메일입니다. 기존 계정을 확인하세요."
    return _render_user_management(request, session, message, single_error)


@router.post("/account-requests/{request_id}/reject")
def reject_account_request(
    request: Request,
    request_id: int,
    session: Session = Depends(get_session),
):
    account_request = session.get(AccountRequest, request_id)
    if not account_request or account_request.status != "pending":
        message, single_error = None, "대기 중인 신청을 찾을 수 없습니다."
    else:
        account_request.status = "rejected"
        account_request.reviewed_at = datetime.utcnow()
        session.commit()
        message, single_error = "계정 신청을 반려했습니다.", None
    return _render_user_management(request, session, message, single_error)


def _render_user_management(
    request: Request, session: Session, message: str | None, single_error: str | None
):
    users = list(
        session.scalars(select(AuthUser).order_by(AuthUser.created_at.desc(), AuthUser.id.desc()))
    )
    requests = list(
        session.scalars(
            select(AccountRequest)
            .where(AccountRequest.status == "pending")
            .order_by(AccountRequest.created_at, AccountRequest.id)
        )
    )
    return templates.TemplateResponse(
        request,
        "admin/users.html",
        {
            "users": users,
            "requests": requests,
            "message": message,
            "errors": [],
            "single_error": single_error,
            "smtp_configured": bool(
                os.getenv("BOM_SMTP_HOST")
                and os.getenv("BOM_SMTP_FROM")
                and (os.getenv("BOM_ADMIN_NOTIFICATION_EMAIL") or users)
            ),
        },
        status_code=status.HTTP_422_UNPROCESSABLE_ENTITY if single_error else status.HTTP_200_OK,
    )


@router.post("/users")
def add_user(
    request: Request,
    email: str = Form(),
    display_name: str = Form(),
    temporary_password: str = Form(),
    session: Session = Depends(get_session),
):
    try:
        user = create_user(
            session,
            email=email,
            display_name=display_name,
            temporary_password=temporary_password,
        )
        session.commit()
        message = f"{user.email} 계정을 등록했습니다. 첫 로그인 때 비밀번호를 변경해야 합니다."
        single_error = None
    except ValueError as error:
        session.rollback()
        message = None
        single_error = str(error)
    except IntegrityError:
        session.rollback()
        message = None
        single_error = "이미 등록된 이메일입니다."
    users = list(
        session.scalars(select(AuthUser).order_by(AuthUser.created_at.desc(), AuthUser.id.desc()))
    )
    return templates.TemplateResponse(
        request,
        "admin/users.html",
        {"users": users, "message": message, "errors": [], "single_error": single_error},
        status_code=status.HTTP_422_UNPROCESSABLE_ENTITY if single_error else status.HTTP_200_OK,
    )


@router.post("/users/import")
async def import_users(
    request: Request,
    account_file: UploadFile = File(),
    session: Session = Depends(get_session),
):
    users = list(
        session.scalars(select(AuthUser).order_by(AuthUser.created_at.desc(), AuthUser.id.desc()))
    )
    if not account_file.filename or not account_file.filename.casefold().endswith(".xlsx"):
        message, errors = None, [".xlsx 형식의 계정 파일을 선택하세요."]
    else:
        contents = await account_file.read(MAX_WORKBOOK_BYTES + 1)
        if len(contents) > MAX_WORKBOOK_BYTES:
            message, errors = None, ["계정 Excel 파일은 2 MB 이하여야 합니다."]
        else:
            accounts, errors = _parse_workbook(contents, session)
            if not errors:
                try:
                    for account in accounts:
                        create_user(session, **account)
                    session.commit()
                    message = (
                        f"{len(accounts)}개 계정을 등록했습니다. "
                        "모든 사용자는 첫 로그인 때 비밀번호를 변경해야 합니다."
                    )
                    users = list(
                        session.scalars(
                            select(AuthUser).order_by(
                                AuthUser.created_at.desc(), AuthUser.id.desc()
                            )
                        )
                    )
                except (ValueError, IntegrityError):
                    session.rollback()
                    message = None
                    errors = ["계정을 저장하지 못했습니다. 중복 여부를 확인하고 다시 업로드하세요."]
                    users = list(session.scalars(select(AuthUser).order_by(AuthUser.id.desc())))
            else:
                message = None
    await account_file.close()
    return templates.TemplateResponse(
        request,
        "admin/users.html",
        {"users": users, "message": message, "errors": errors, "single_error": None},
        status_code=status.HTTP_422_UNPROCESSABLE_ENTITY if errors else status.HTTP_200_OK,
    )
