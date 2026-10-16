"""Excel 출력의 범위·금액·고객 정보 분리 회귀 테스트."""

from decimal import Decimal
from pathlib import Path
from types import SimpleNamespace

import pytest
from openpyxl import Workbook, load_workbook

from app.exporters.excel import assert_customer_safe, export_cost_summary, export_customer_quote


def _quote_item(**overrides: object) -> SimpleNamespace:
    values: dict[str, object] = {
        "drawing_number": "DWG-001",
        "item_name": "검증 품목",
        "item_type": "purchased",
        "material": None,
        "treatment": None,
        "specification": "M10",
        "maker": "제조사",
        "quantity": Decimal("2"),
        "material_cost": Decimal("0"),
        "machining_cost": Decimal("0"),
        "treatment_cost": Decimal("0"),
        "purchased_unit_price": Decimal("1200"),
    }
    values.update(overrides)
    return SimpleNamespace(**values)


def _customer_template(path: Path) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "2.Detail Quotation(Global)"
    sheet["B5"] = "Purchasing Item"
    sheet["C6"] = "1)"
    sheet["C7"] = "Purchasing Item Sub Total"
    sheet["B8"] = "Processing Item"
    sheet["C9"] = "1)"
    sheet["C10"] = "2)"
    sheet["C11"] = "3)"
    workbook.save(path)


def test_customer_quote_expands_machined_costs_without_overwriting_totals(tmp_path: Path) -> None:
    template = tmp_path / "template.xlsx"
    output = tmp_path / "customer.xlsx"
    _customer_template(template)
    quote = SimpleNamespace(version_number=2)
    items = [
        _quote_item(),
        _quote_item(
            item_type="machined",
            item_name="가공 브라켓",
            material="AL6061",
            treatment="아노다이징",
            material_cost=Decimal("100"),
            machining_cost=Decimal("200"),
            treatment_cost=Decimal("30"),
            purchased_unit_price=Decimal("0"),
        ),
    ]

    export_customer_quote(template, output, quote=quote, items=items)

    sheet = load_workbook(output, data_only=False)["2.Detail Quotation(Global)"]
    assert sheet["A2"].value == "Quotation 2"
    assert (sheet["D6"].value, sheet["E6"].value, sheet["I6"].value) == (
        "검증 품목",
        "M10",
        1200.0,
    )
    assert [sheet.cell(row, 9).value for row in (9, 10, 11)] == [100.0, 200.0, 30.0]
    assert sheet["C7"].value == "Purchasing Item Sub Total"
    assert sheet.parent.calculation.fullCalcOnLoad is True


def test_customer_output_rejects_internal_fields() -> None:
    with pytest.raises(ValueError, match="금지 필드"):
        assert_customer_safe([{"drawing_number": "DWG-001", "supplier": "내부 거래처"}])


def test_cost_summary_uses_customer_selling_components_only(tmp_path: Path) -> None:
    output = tmp_path / "cost-summary.xlsx"
    export_cost_summary(
        output,
        items=[
            _quote_item(
                item_type="machined",
                material_cost=Decimal("100"),
                machining_cost=Decimal("200"),
                treatment_cost=Decimal("30"),
            )
        ],
    )

    sheet = load_workbook(output, data_only=True).active
    assert sheet["H2"].value == 660.0
